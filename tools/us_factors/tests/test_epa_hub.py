"""
The EPA Hub parser: declared tables, token-matched columns, unit normalisation.

Values in the fixture are invented round numbers -- see ``conftest.py`` for why.
These tests are about whether the right cell reaches the right row.
"""
# ruff: noqa: TC003  # CLAUDE.md forbids `from __future__ import annotations`
# in new files, and without it the flake8-type-checking fixes are not available.



from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook

from tools.us_factors.sources import epa_hub
from tools.us_factors.sources.workbook import WorkbookLayoutError
from tools.us_factors.tidy import validate_long

SOURCE = 'US EPA GHG Emission Factors Hub 2025'


def test_parses_stationary_combustion(epa_hub_workbook: Path) -> None:
    df = epa_hub.parse_table(epa_hub_workbook, epa_hub.STATIONARY, SOURCE, data_year=2025)

    assert set(df['fuel'].unique()) == {'natural_gas', 'propane'}
    assert set(df['greenhouse_gas'].unique()) == {'co2', 'ch4', 'n2o'}
    validate_long(df)


def test_slugifies_published_labels_into_category_ids(epa_hub_workbook: Path) -> None:
    df = epa_hub.parse_table(epa_hub_workbook, epa_hub.MOBILE_MILEAGE, SOURCE, data_year=2025)

    assert set(df['vehicle_type'].unique()) == {'passenger_car', 'light_duty_truck'}


def test_converts_gram_columns_into_the_declared_kilogram_unit(epa_hub_workbook: Path) -> None:
    """CH4/N2O are published in g/MMBtu; the metric is kg/MMBtu."""
    df = epa_hub.parse_table(epa_hub_workbook, epa_hub.STATIONARY, SOURCE, data_year=2025)

    assert df['unit'].unique().to_list() == ['kg/MMBtu']
    ch4_gas = df.filter((pl.col('fuel') == 'natural_gas') & (pl.col('greenhouse_gas') == 'ch4'))
    # Fixture publishes 1.0 g/MMBtu -> 0.001 kg/MMBtu.
    assert ch4_gas['value'].item() == pytest.approx(0.001)
    assert 'published in g/MMBtu' in ch4_gas['comment'].item()


def test_co2_column_is_not_rescaled(epa_hub_workbook: Path) -> None:
    df = epa_hub.parse_table(epa_hub_workbook, epa_hub.STATIONARY, SOURCE, data_year=2025)

    co2 = df.filter((pl.col('fuel') == 'natural_gas') & (pl.col('greenhouse_gas') == 'co2'))
    assert co2['value'].item() == pytest.approx(50.0)


def test_token_matching_survives_reworded_headers(tmp_path: Path) -> None:
    """A rewritten header with the same tokens must still be found."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Stationary Combustion'
    sheet.append(['Fuel type', 'CO2 Factor (kg / mmBtu)', 'CH4 Factor (g / mmBtu)'])
    sheet.append(['Natural Gas', 50.0, 1.0])
    path = tmp_path / 'reworded.xlsx'
    wb.save(path)

    df = epa_hub.parse_table(path, epa_hub.STATIONARY, SOURCE, data_year=2025)

    assert df.filter(pl.col('greenhouse_gas') == 'co2')['value'].item() == pytest.approx(50.0)


def test_does_not_confuse_one_gas_column_for_another(tmp_path: Path) -> None:
    """The CO2 tokens must not match the CH4 column when CO2 is absent."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Stationary Combustion'
    sheet.append(['Fuel type', 'g CH4 per mmBtu'])
    sheet.append(['Natural Gas', 1.0])
    path = tmp_path / 'ch4_only.xlsx'
    wb.save(path)

    df = epa_hub.parse_table(path, epa_hub.STATIONARY, SOURCE, data_year=2025)

    assert df['greenhouse_gas'].unique().to_list() == ['ch4']


def test_refuses_when_no_declared_column_matches(tmp_path: Path) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Stationary Combustion'
    sheet.append(['Fuel type', 'Something Unrelated'])
    sheet.append(['Natural Gas', 1.0])
    path = tmp_path / 'no_match.xlsx'
    wb.save(path)

    with pytest.raises(WorkbookLayoutError, match='none of the declared value columns'):
        epa_hub.parse_table(path, epa_hub.STATIONARY, SOURCE, data_year=2025)


def test_refuses_an_ambiguous_column_rather_than_picking_one(tmp_path: Path) -> None:
    """Two columns matching the same tokens is a layout change, not a coin flip."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Stationary Combustion'
    sheet.append(['Fuel type', 'kg CO2 per mmBtu', 'kg CO2 per mmBtu (biogenic)'])
    sheet.append(['Natural Gas', 50.0, 0.0])
    path = tmp_path / 'ambiguous.xlsx'
    wb.save(path)

    with pytest.raises(WorkbookLayoutError, match='matched several columns'):
        epa_hub.parse_table(path, epa_hub.STATIONARY, SOURCE, data_year=2025)


def test_stops_at_the_blank_row_between_stacked_tables(epa_hub_workbook: Path) -> None:
    """Hub sheets stack tables; only the first belongs to this declaration."""
    df = epa_hub.parse_table(epa_hub_workbook, epa_hub.STATIONARY, SOURCE, data_year=2025)

    assert df['fuel'].n_unique() == 2
