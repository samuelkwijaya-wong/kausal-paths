"""
The eGRID parser: finding the right cells, and refusing when it cannot.

The refusal tests matter as much as the happy path. A parser that quietly returns
fewer rows when a column moves is how a city ends up with a missing grid factor
and a model that still computes.
"""
# ruff: noqa: TC003  # CLAUDE.md forbids `from __future__ import annotations`
# in new files, and without it the flake8-type-checking fixes are not available.



from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook, load_workbook

from tools.us_factors.sources import egrid
from tools.us_factors.sources.workbook import WorkbookLayoutError
from tools.us_factors.tidy import validate_long

SOURCE = 'US EPA eGRID2023'


def test_parses_every_subregion_and_gas(egrid_workbook: Path) -> None:
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    assert set(df['egrid_subregion'].unique()) == {'AZNM', 'CAMX', 'NEWE'}
    assert set(df['greenhouse_gas'].unique()) == {'co2', 'ch4', 'n2o', 'co2e'}
    assert df.height == 3 * 4


def test_reads_values_against_the_right_subregion(egrid_workbook: Path) -> None:
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    co2 = df.filter(pl.col('greenhouse_gas') == 'co2').sort('egrid_subregion')
    assert co2.select(['egrid_subregion', 'value']).rows() == [
        ('AZNM', 700.0),
        ('CAMX', 500.0),
        ('NEWE', 600.0),
    ]


def test_normalises_per_gwh_rates_to_one_metric_unit(egrid_workbook: Path) -> None:
    """
    Per-GWh rates are converted to the metric unit.

    EPA publishes CH4/N2O per GWh and CO2 per MWh, but the upload format allows
    one unit per metric -- so the per-GWh rates are converted on the way in.
    """
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    assert df['unit'].unique().to_list() == ['lb/MWh']

    ch4_aznm = df.filter(
        (pl.col('greenhouse_gas') == 'ch4') & (pl.col('egrid_subregion') == 'AZNM')
    )
    # The fixture publishes 50.0 lb/GWh, which is 0.05 lb/MWh.
    assert ch4_aznm['value'].item() == pytest.approx(0.05)


def test_records_the_conversion_so_a_value_stays_checkable(egrid_workbook: Path) -> None:
    """A converted value must say what it was published as, or it cannot be audited."""
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    ch4_comment = df.filter(pl.col('greenhouse_gas') == 'ch4')['comment'].unique().item()
    assert 'lb/GWh' in ch4_comment
    assert 'SRCH4RTA' in ch4_comment

    co2_comment = df.filter(pl.col('greenhouse_gas') == 'co2')['comment'].unique().item()
    assert 'published as' not in co2_comment


def test_labels_rows_with_the_data_year_not_the_release(egrid_workbook: Path) -> None:
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2021, source_label=SOURCE)

    assert df['year'].unique().to_list() == [2021]
    assert df['source'].unique().to_list() == [SOURCE]


def test_skips_footnote_rows_and_the_descriptive_row(egrid_workbook: Path) -> None:
    """Neither the 'ZZZZ' footnote nor the descriptive row may become a factor."""
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    assert 'ZZZZ' not in df['egrid_subregion'].to_list()
    assert df['value'].dtype == pl.Float64


def test_ignores_the_state_sheet(egrid_workbook: Path) -> None:
    """State rates answer a different question and must not leak into the subregion table."""
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    assert 999.0 not in df['value'].to_list()


def test_output_satisfies_the_long_frame_contract(egrid_workbook: Path) -> None:
    df = egrid.parse(egrid_workbook, release_year=2023, data_year=2023, source_label=SOURCE)

    validate_long(df)


def test_finds_the_sheet_when_the_release_suffix_moves(egrid_workbook: Path, tmp_path: Path) -> None:
    """A release that names its sheet differently must still parse."""
    wb = load_workbook(egrid_workbook)
    wb['SRL23'].title = 'SRL24'
    renamed = tmp_path / 'egrid_renamed.xlsx'
    wb.save(renamed)

    df = egrid.parse(renamed, release_year=2024, data_year=2024, source_label=SOURCE)

    assert df.height == 3 * 4


def test_refuses_a_workbook_with_no_subregion_sheet(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.title = 'Notes'
    wb.active.append(['nothing here'])
    path = tmp_path / 'empty.xlsx'
    wb.save(path)

    with pytest.raises(WorkbookLayoutError, match='subregion sheet'):
        egrid.parse(path, release_year=2023, data_year=2023, source_label=SOURCE)


def test_refuses_when_every_rate_column_is_missing(tmp_path: Path) -> None:
    """A renamed rate column must stop the build, not silently yield nothing."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'SRL23'
    sheet.append(['SUBRGN', 'SRNAME', 'SOMETHING_ELSE'])
    sheet.append(['AZNM', 'WECC Southwest', 1.0])
    path = tmp_path / 'no_rates.xlsx'
    wb.save(path)

    with pytest.raises(WorkbookLayoutError, match='none of the expected rate columns'):
        egrid.parse(path, release_year=2023, data_year=2023, source_label=SOURCE)


def test_refuses_an_unknown_release_rather_than_guessing_a_url() -> None:
    with pytest.raises(WorkbookLayoutError, match='No download URL recorded'):
        egrid.spec_for_release(1999)
