"""
Locating data inside a published workbook without hard-coding cell positions.

EPA and NREL re-lay-out their spreadsheets between releases: a banner row is
added, a column moves, a sheet is renamed from ``SRL22`` to ``SRL23``. A parser
written against fixed coordinates reads the wrong column silently, which is the
one failure mode this package cannot tolerate.

So every lookup here is by *content*: find the sheet whose name matches a
pattern, find the header row by the anchor token it must contain, find a column
by its header text. When a lookup fails it raises with what it searched for and
what it found instead, so a layout change is a five-minute fix rather than an
investigation.
"""
# ruff: noqa: TC002, TC003  # CLAUDE.md forbids `from __future__ import annotations`
# in new files, and without it the flake8-type-checking fixes are not available.


import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

MAX_HEADER_SCAN_ROWS = 25
"""How far down to look for the header row. Real workbooks put banners above it."""


class WorkbookLayoutError(RuntimeError):
    """A workbook did not contain what the parser expected, with the detail needed to fix it."""


@dataclass
class HeaderedSheet:
    """A worksheet plus the located header row, addressable by column name."""

    sheet: Worksheet
    header_row: int
    columns: dict[str, int]
    """Normalised header text -> 1-based column index."""

    @property
    def title(self) -> str:
        return self.sheet.title

    def has(self, *names: str) -> bool:
        return all(normalise(n) in self.columns for n in names)

    def column_index(self, name: str) -> int:
        key = normalise(name)
        if key not in self.columns:
            close = [c for c in self.columns if key in c or c in key]
            raise WorkbookLayoutError(
                f"Sheet '{self.title}' has no column '{name}'.\n"
                f'  Header row: {self.header_row}\n'
                f'  Near matches: {close or "none"}\n'
                f'  Columns present: {sorted(self.columns)[:40]}'
            )
        return self.columns[key]

    def find_by_tokens(self, tokens: tuple[str, ...]) -> str | None:
        """
        Return the single column whose header contains every token, or None.

        Publishers write a header as prose -- ``kg CO2 per mmBtu``, ``CO2 Factor
        (kg / mmBtu)`` -- so matching on the tokens that carry the meaning
        survives rewording that an exact name would not. Ambiguity is an error
        rather than a first-match, because picking the wrong one of two CO2
        columns is silent.
        """
        wanted = [normalise(t) for t in tokens]
        hits = [name for name in self.columns if all(t in name for t in wanted)]
        if not hits:
            return None
        if len(hits) > 1:
            raise WorkbookLayoutError(
                f"Tokens {tokens} matched several columns in '{self.title}': {hits}. "
                'Add a distinguishing token.'
            )
        return hits[0]

    def rows(self) -> list[dict[str, Any]]:
        """Return the data rows below the header, each keyed by normalised column name."""
        by_index = {idx: name for name, idx in self.columns.items()}
        out: list[dict[str, Any]] = []
        for row in self.sheet.iter_rows(min_row=self.header_row + 1):
            record = {
                by_index[idx]: cell.value
                for idx, cell in enumerate(row, start=1)
                if idx in by_index
            }
            if any(v is not None for v in record.values()):
                out.append(record)
        return out


def normalise(text: object) -> str:
    """Fold header text to a stable key: lowercase, no punctuation or runs of space."""
    return re.sub(r'[^a-z0-9]+', ' ', str(text).lower()).strip()


def open_workbook(path: Path) -> Any:
    """Open read-only with formulas evaluated, which is what a published workbook carries."""
    return load_workbook(path, data_only=True, read_only=True)


def find_sheet(workbook: Any, pattern: str) -> Worksheet:
    """
    Return the single sheet whose name matches ``pattern`` (a case-insensitive regex).

    Matching by pattern rather than exact name is what lets one parser read
    ``SRL21`` through ``SRL23`` without a table of per-release sheet names.
    """
    rx = re.compile(pattern, re.IGNORECASE)
    matches = [name for name in workbook.sheetnames if rx.search(name)]
    if not matches:
        raise WorkbookLayoutError(
            f"No sheet matching /{pattern}/.\n  Sheets present: {workbook.sheetnames}"
        )
    if len(matches) > 1:
        raise WorkbookLayoutError(
            f"Pattern /{pattern}/ matched several sheets: {matches}. Tighten the pattern."
        )
    return workbook[matches[0]]


def read_headered(sheet: Worksheet, anchor: str) -> HeaderedSheet:
    """
    Locate the header row as the first row containing ``anchor``, and index it.

    ``anchor`` is a column name the sheet must have (``SUBRGN`` for an eGRID
    subregion sheet). Scanning for it skips whatever banner rows the publisher
    added this year.
    """
    wanted = normalise(anchor)
    for row_idx, row in enumerate(sheet.iter_rows(max_row=MAX_HEADER_SCAN_ROWS), start=1):
        cells = {
            normalise(cell.value): idx
            for idx, cell in enumerate(row, start=1)
            if cell.value is not None
        }
        if wanted in cells:
            return HeaderedSheet(sheet=sheet, header_row=row_idx, columns=cells)

    seen = [
        [c.value for c in row[:12]]
        for row in sheet.iter_rows(max_row=min(5, MAX_HEADER_SCAN_ROWS))
    ]
    raise WorkbookLayoutError(
        f"Sheet '{sheet.title}' has no header row containing '{anchor}' "
        f'in its first {MAX_HEADER_SCAN_ROWS} rows.\n  First rows were: {seen}'
    )


def describe(path: Path, max_cols: int = 30) -> str:
    """
    Render a human-readable dump of a workbook's structure.

    This is what ``build_csv.py --inspect`` prints. It exists because the person
    running the build against a new release needs to see what changed without
    opening Excel, and because a parser fix starts from exactly this information.
    """
    workbook = open_workbook(path)
    lines = [f'{path.name}', f'  sheets ({len(workbook.sheetnames)}): {workbook.sheetnames}', '']
    for name in workbook.sheetnames:
        sheet = workbook[name]
        lines.append(f'  [{name}] dims={sheet.calculate_dimension()}')
        for row_idx, row in enumerate(sheet.iter_rows(max_row=4), start=1):
            values = [c.value for c in row[:max_cols]]
            if any(v is not None for v in values):
                lines.append(f'    row {row_idx}: {values}')
        lines.append('')
    return '\n'.join(lines)
